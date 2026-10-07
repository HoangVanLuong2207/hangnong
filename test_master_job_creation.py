from __future__ import annotations

import json
import io
import sqlite3
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
from unittest import mock
from pathlib import Path

import master_server
from master_server import (
    CoordinatorServer,
    LocalStore,
    MasterHandler,
    _prune_completed_jobs_before_today,
    _today_start_timestamp,
)


class _BlockingCreateStore:
    """Pause the create batch after the job row exists but before chunks exist."""

    def __init__(self, inner: LocalStore) -> None:
        self.inner = inner
        self.batch_entered = threading.Event()
        self.release_batch = threading.Event()
        self.block_once = True

    def __getattr__(self, name: str):
        return getattr(self.inner, name)

    def batch(self, statements: list) -> None:
        opens_job = any(
            isinstance(statement, dict)
            and "UPDATE jobs SET status='open'" in statement.get("sql", "")
            for statement in statements
        )
        if opens_job and self.block_once:
            self.block_once = False
            self.batch_entered.set()
            if not self.release_batch.wait(5):
                raise TimeoutError("test did not release create batch")
        self.inner.batch(statements)


class _CountingStore:
    """Record read calls while delegating storage to LocalStore."""

    def __init__(self, inner: LocalStore) -> None:
        self.inner = inner
        self.read_calls = 0

    def __getattr__(self, name: str):
        return getattr(self.inner, name)

    def fetch(self, sql: str, args: tuple = ()) -> list[tuple]:
        self.read_calls += 1
        return self.inner.fetch(sql, args)

    def fetchone(self, sql: str, args: tuple = ()) -> tuple | None:
        self.read_calls += 1
        return self.inner.fetchone(sql, args)


class ParseAccountsValidationTest(unittest.TestCase):
    def test_format_error_includes_line_number_and_original_content(self) -> None:
        with self.assertRaises(ValueError) as raised:
            master_server.parse_accounts("valid|password\n  invalid data  ")

        self.assertEqual(
            str(raised.exception),
            'Dòng 2: cần định dạng user|pass, user|pass|mail hoặc '
            'user|pass|mail|passmail (hoặc user:pass). '
            'Nội dung dòng: "  invalid data  "',
        )

    def test_invalid_credentials_error_includes_line_and_content(self) -> None:
        with self.assertRaises(ValueError) as raised:
            master_server.parse_accounts("valid|password\n|missing-user")

        self.assertEqual(
            str(raised.exception),
            'Dòng 2: tài khoản/mật khẩu không hợp lệ. '
            'Nội dung dòng: "|missing-user"',
        )

    def test_error_content_keeps_unicode_readable(self) -> None:
        with self.assertRaises(ValueError) as raised:
            master_server.parse_accounts("tài khoản không có dấu phân cách")

        self.assertIn(
            'Nội dung dòng: "tài khoản không có dấu phân cách"',
            str(raised.exception),
        )

    def test_username_rules_partition_valid_and_invalid_accounts(self) -> None:
        accounts = master_server.parse_accounts(
            "Valid_1|pass\na-b.c9|pass\n_badusr|pass\n123456|pass\nabc12|pass\nabc@123|pass"
        )

        valid, invalid = master_server.partition_accounts_by_username(accounts)

        self.assertEqual([item.account for item in valid], ["Valid_1", "a-b.c9"])
        self.assertEqual(
            [item.username for item in invalid],
            ["_badusr", "123456", "abc12", "abc@123"],
        )
        self.assertEqual([item.line_number for item in invalid], [3, 4, 5, 6])

    def test_junk_rules_cover_placeholders_urls_tokens_profiles_and_weak_passwords(self) -> None:
        accounts = master_server.parse_accounts("\n".join((
            "wirawan0799|UNKNOWN",
            "maps.google.com|Secure123!",
            "profile01|Google_[Chrome]_Profile4",
            "tokenusr|c2b598c2-f0b7-4f45-9cab-ed5c19ed9177",
            "shortusr|234",
            "nospecial|Secure123",
            "jesran|jesran",
            "В±Р’В»Р’В§y|Thaithuy96",
            "gooduser|Secure123!",
        )))

        valid, invalid, junk = master_server.partition_accounts_for_job(accounts)

        self.assertEqual([item.account for item in valid], ["nospecial", "jesran", "gooduser"])
        self.assertEqual(invalid, [])
        self.assertEqual(len(junk), 6)
        self.assertIn("Placeholder", junk[0].reason)
        self.assertIn("domain", junk[1].reason)
        self.assertIn("Chrome profile", junk[2].reason)
        self.assertIn("UUID", junk[3].reason)
        self.assertIn("ngắn hơn 6", junk[4].reason)
        self.assertIn("mojibake", junk[5].reason)

    def test_legacy_password_without_special_character_is_allowed(self) -> None:
        self.assertIsNone(master_server.junk_account_reason(
            master_server.ParsedAccount("gooduser", "Secure123!"),
        ))
        self.assertIsNone(master_server.junk_account_reason(
            master_server.ParsedAccount("nospecial", "Secure123"),
        ))
        self.assertIsNone(master_server.junk_account_reason(
            master_server.ParsedAccount("sameuser", "sameuser"),
        ))
        self.assertIsNone(master_server.junk_account_reason(
            master_server.ParsedAccount("sameuser", "sameuser..."),
        ))
        self.assertIsNone(master_server.junk_account_reason(
            master_server.ParsedAccount("dotpass1", "Secure.xyz"),
        ))
        self.assertIsNone(master_server.junk_account_reason(
            master_server.ParsedAccount("single.com", "Secure123!"),
        ))
        self.assertIsNone(master_server.junk_account_reason(
            master_server.ParsedAccount("gooduser", "sso.garena.com"),
        ))
        self.assertIn("domain", master_server.junk_account_reason(
            master_server.ParsedAccount("maps.google.com", "Secure123!"),
        ) or "")

    def test_repeated_users_and_passwords_are_allowed(self) -> None:
        text = "\n".join(
            [f"groupa{index}|danan777!" for index in range(10)]
            + [f"groupb{index}|danan777!" for index in range(10)]
            + ["sameuser|Alpha123!", "sameuser|Beta456!"]
        )

        valid, invalid, junk = master_server.partition_accounts_for_job(
            master_server.parse_accounts(text)
        )

        self.assertEqual(len(valid), 22)
        self.assertEqual([item.account for item in valid[-2:]], ["sameuser", "sameuser"])
        self.assertEqual(invalid, [])
        self.assertEqual(junk, [])


class InvalidUsernameJobFlowTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.store = LocalStore(Path(self.temp_dir.name) / "invalid-user-test.db")
        self.server = CoordinatorServer(("127.0.0.1", 0), MasterHandler, self.store, "secret")
        self.server_thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.server_thread.start()
        self.base_url = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.server_thread.join(5)
        self.store._conn.close()
        self.temp_dir.cleanup()

    def post_error(self, body: dict) -> tuple[int, dict]:
        request = urllib.request.Request(
            self.base_url + "/api/jobs",
            data=json.dumps(body).encode("utf-8"),
            method="POST",
            headers={"Authorization": "Bearer secret", "Content-Type": "application/json"},
        )
        try:
            with urllib.request.urlopen(request, timeout=10) as response:
                return response.status, json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            try:
                return exc.code, json.loads(exc.read().decode("utf-8"))
            finally:
                exc.close()

    def get_json(self, path: str) -> dict:
        request = urllib.request.Request(
            self.base_url + path,
            headers={"Authorization": "Bearer secret"},
        )
        with urllib.request.urlopen(request, timeout=10) as response:
            return json.loads(response.read().decode("utf-8"))

    def test_warns_then_filters_stores_groups_and_exports_rejected_accounts(self) -> None:
        body = {
            "text": "Valid_1|secret12!\n_badusr|secret12!\nvalidtwo|UNKNOWN",
            "billing_mode": "quantity",
        }
        status, warning = self.post_error(body)

        self.assertEqual(status, 409)
        self.assertEqual(warning["code"], "FILTERED_ACCOUNTS_REQUIRE_CONFIRMATION")
        self.assertEqual(warning["valid_accounts"], 1)
        self.assertEqual(warning["invalid_accounts"], 1)
        self.assertEqual(warning["junk_accounts_count"], 1)
        self.assertEqual(self.store.fetchone("SELECT COUNT(*) FROM jobs")[0], 0)

        body["confirm_filtered_accounts"] = True
        status, created = self.post_error(body)
        self.assertEqual(status, 200)
        self.assertTrue(created["ok"])
        self.assertEqual(created["total"], 1)
        self.assertEqual(created["invalid_accounts"], 1)
        self.assertEqual(created["junk_accounts_count"], 1)
        job_id = created["job_id"]

        saved = self.store.fetchone(
            "SELECT line_number,username,raw_line,reason FROM invalid_usernames WHERE job_id=?",
            (job_id,),
        )
        self.assertEqual(saved[:3], (2, "_badusr", "_badusr|secret12!"))
        self.assertIn("bắt đầu", saved[3])

        invalid_group = self.get_json(f"/api/jobs/{job_id}/rows?filter=INVALID")
        self.assertEqual(invalid_group["category_counts"]["INVALID"], 1)
        self.assertEqual(invalid_group["rows"][0]["account"], "_badusr")

        junk_group = self.get_json(f"/api/jobs/{job_id}/rows?filter=JUNK")
        self.assertEqual(junk_group["category_counts"]["JUNK"], 1)
        self.assertEqual(junk_group["rows"][0]["account"], "validtwo")
        all_rejected = self.get_json(f"/api/jobs/{job_id}/rows?filter=ALL")
        self.assertEqual(all_rejected["total_all"], 2)
        self.assertEqual({row["result_type"] for row in all_rejected["rows"]}, {
            "invalid username", "junk data",
        })

        request = urllib.request.Request(
            self.base_url + f"/api/jobs/{job_id}/export.xlsx?min_level=12",
            headers={"Authorization": "Bearer secret"},
        )
        with urllib.request.urlopen(request, timeout=10) as response:
            workbook_bytes = response.read()
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True)
        try:
            self.assertIn("USER không hợp lệ", workbook.sheetnames)
            self.assertIn("DATA rác", workbook.sheetnames)
            row = list(workbook["USER không hợp lệ"].iter_rows(min_row=2, values_only=True))[0]
            self.assertEqual(row[0], "2")
            self.assertEqual(row[1], "_badusr|secret12!")
            junk_row = list(workbook["DATA rác"].iter_rows(min_row=2, values_only=True))[0]
            self.assertEqual(junk_row[0], "3")
            self.assertEqual(junk_row[1], "validtwo|UNKNOWN")
        finally:
            workbook.close()


class PostgreSQLMigrationOrderTest(unittest.TestCase):
    def test_legacy_columns_are_migrated_before_schema_indexes(self) -> None:
        table_statements, index_statements = master_server._postgres_schema_phases()

        self.assertTrue(table_statements)
        self.assertTrue(index_statements)
        self.assertFalse(any(statement.upper().startswith("CREATE INDEX") for statement in table_statements))
        self.assertFalse(any(statement.upper().startswith("CREATE UNIQUE INDEX") for statement in table_statements))
        self.assertTrue(any("jobs(owner_user_id)" in statement for statement in index_statements))
        self.assertTrue(any("jobs(external_job_reference)" in statement for statement in index_statements))


class JobCreationRaceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.inner = LocalStore(Path(self.temp_dir.name) / "master-test.db")
        self.store = _BlockingCreateStore(self.inner)
        self.server = CoordinatorServer(("127.0.0.1", 0), MasterHandler, self.store, "secret")
        self.server_thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.server_thread.start()
        self.base_url = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self) -> None:
        self.store.release_batch.set()
        self.server.shutdown()
        self.server.server_close()
        self.server_thread.join(5)
        self.inner._conn.close()
        self.temp_dir.cleanup()

    def post(self, path: str, body: dict, token: str = "secret") -> tuple[int, dict]:
        request = urllib.request.Request(
            self.base_url + path,
            data=json.dumps(body).encode("utf-8"),
            method="POST",
            headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        )
        with urllib.request.urlopen(request, timeout=10) as response:
            return response.status, json.loads(response.read().decode("utf-8"))

    def post_error(self, path: str, body: dict, token: str = "secret") -> tuple[int, dict]:
        try:
            return self.post(path, body, token)
        except urllib.error.HTTPError as exc:
            try:
                return exc.code, json.loads(exc.read().decode("utf-8"))
            finally:
                exc.close()

    def get(self, path: str, token: str = "secret") -> tuple[int, dict]:
        request = urllib.request.Request(
            self.base_url + path,
            headers={"Authorization": f"Bearer {token}"},
        )
        with urllib.request.urlopen(request, timeout=10) as response:
            return response.status, json.loads(response.read().decode("utf-8"))

    def get_error(self, path: str, token: str = "secret") -> tuple[int, dict]:
        try:
            return self.get(path, token)
        except urllib.error.HTTPError as exc:
            try:
                return exc.code, json.loads(exc.read().decode("utf-8"))
            finally:
                exc.close()

    def wait_for_job_status(
        self, job_id: int, expected: str = "done", timeout: float = 5.0
    ) -> str:
        deadline = time.monotonic() + timeout
        status = ""
        while time.monotonic() < deadline:
            row = self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (job_id,))
            status = str(row[0]) if row else "missing"
            if status == expected:
                return status
            time.sleep(0.02)
        self.fail(f"job {job_id} status={status!r}, expected {expected!r}")

    def test_create_job_format_error_reports_line_content(self) -> None:
        status, payload = self.post_error(
            "/api/jobs",
            {"text": "valid|password!\ninvalid data"},
        )

        self.assertEqual(status, 400)
        self.assertFalse(payload["ok"])
        self.assertIn("Dòng 2:", payload["error"])
        self.assertIn('Nội dung dòng: "invalid data"', payload["error"])

    def test_claim_cannot_finish_job_before_chunks_are_saved(self) -> None:
        created: dict = {}

        def create_job() -> None:
            try:
                text = "\n".join(f"user{index:02d}|password{index}!" for index in range(31))
                created["response"] = self.post("/api/jobs", {"text": text, "billing_mode": "quantity"})
            except Exception as exc:  # pragma: no cover - surfaced by assertion below
                created["error"] = exc

        create_thread = threading.Thread(target=create_job)
        create_thread.start()
        self.assertTrue(self.store.batch_entered.wait(5))

        job = self.inner.fetchone("SELECT id, status FROM jobs ORDER BY id DESC LIMIT 1")
        self.assertIsNotNone(job)
        self.assertEqual(job[1], "creating")

        status, claim = self.post("/api/claim", {"satellite_id": "race-satellite"})
        self.assertEqual(status, 200)
        self.assertIsNone(claim["claim"])
        self.assertEqual(
            self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (job[0],))[0],
            "creating",
        )

        self.store.release_batch.set()
        create_thread.join(10)
        self.assertFalse(create_thread.is_alive())
        self.assertNotIn("error", created)
        self.assertEqual(created["response"][0], 200)

        job_id = created["response"][1]["job_id"]
        self.assertEqual(self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (job_id,))[0], "open")
        self.assertEqual(self.inner.fetchone("SELECT COUNT(*) FROM chunks WHERE job_id=?", (job_id,))[0], 3)

        status, claim = self.post("/api/claim", {"satellite_id": "race-satellite"})
        self.assertEqual(status, 200)
        self.assertEqual(claim["claim"]["job_id"], job_id)

    def test_same_username_with_multiple_passwords_keeps_every_result(self) -> None:
        self.store.block_once = False
        status, created = self.post(
            "/api/jobs",
            {"text": "sameuser|Alpha123!\nsameuser|Beta456!", "billing_mode": "quantity"},
        )
        self.assertEqual(status, 200)
        job_id = created["job_id"]
        chunk_id = self.inner.fetchone(
            "SELECT id FROM chunks WHERE job_id=? ORDER BY idx LIMIT 1", (job_id,)
        )[0]

        status, reported = self.post(
            "/api/report",
            {
                "chunk_id": chunk_id,
                "done": True,
                "rows": [
                    {"stt": "1", "account": "sameuser", "status": "FAIL", "result_type": "Sai pass"},
                    {"stt": "2", "account": "sameuser", "status": "OK", "result_type": "Đúng pass"},
                ],
            },
        )

        self.assertEqual(status, 200)
        self.assertTrue(reported["ok"])
        rows = self.inner.fetch(
            "SELECT account,row_json FROM results WHERE job_id=? ORDER BY id", (job_id,)
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual([json.loads(row_json)["stt"] for _, row_json in rows], ["1", "2"])
        self.assertNotEqual(rows[0][0], rows[1][0])

    def test_normal_and_vvip_satellites_claim_only_their_own_queue(self) -> None:
        normal_job = self.inner.exec(
            "INSERT INTO jobs (created_at,total,chunk_size,status,queue_type) VALUES (?,?,?,?,?)",
            (1, 1, 15, "open", "normal"),
        )
        vvip_job = self.inner.exec(
            "INSERT INTO jobs (created_at,total,chunk_size,status,queue_type) VALUES (?,?,?,?,?)",
            (2, 1, 15, "open", "vvip"),
        )
        self.inner.exec("INSERT INTO chunks (job_id,idx,account) VALUES (?,?,?)", (normal_job, 0, '["normal|pass"]'))
        self.inner.exec("INSERT INTO chunks (job_id,idx,account) VALUES (?,?,?)", (vvip_job, 0, '["vvip|pass"]'))

        status, normal_claim = self.post("/api/claim", {"satellite_id": "normal-satellite"})
        self.assertEqual(status, 200)
        self.assertEqual(normal_claim["claim"]["job_id"], normal_job)
        self.assertEqual(normal_claim["claim"]["queue_type"], "normal")

        denied_status, _ = self.post_error(
            "/api/vvip/claim", {"satellite_id": "fake-vvip"}, token="wrong-secret"
        )
        self.assertEqual(denied_status, 401)

        status, vvip_claim = self.post(
            "/api/vvip/claim", {"satellite_id": "vvip-satellite"}, token="secret"
        )
        self.assertEqual(status, 200)
        self.assertEqual(vvip_claim["claim"]["job_id"], vvip_job)
        self.assertEqual(vvip_claim["claim"]["queue_type"], "vvip")

    def test_admin_defaults_to_vvip_queue_job(self) -> None:
        self.store.block_once = False
        status, created = self.post(
            "/api/jobs",
            {"text": "vvip-user|password!"},
        )
        self.assertEqual(status, 200)
        self.assertEqual(created["queue_type"], "vvip")
        job = self.inner.fetchone(
            "SELECT billing_mode,queue_type FROM jobs WHERE id=?", (created["job_id"],)
        )
        self.assertEqual(job, ("admin", "vvip"))

        _, vvip_claim = self.post(
            "/api/vvip/claim", {"satellite_id": "vvip-satellite"}, token="secret"
        )
        self.assertEqual(vvip_claim["claim"]["job_id"], created["job_id"])

    def test_vvip_chunk_shares_are_30_70_including_small_jobs(self) -> None:
        for count in (1, 2, 3, 4, 7, 9, 10, 11, 20, 23):
            with self.subTest(chunks=count):
                job_id = self.inner.exec(
                    "INSERT INTO jobs (created_at,total,chunk_size,status,queue_type) VALUES (?,?,?,?,?)",
                    (time.time(), count, 1, "open", "vvip"),
                )
                for idx in range(count):
                    self.inner.exec(
                        "INSERT INTO chunks (job_id,idx,account) VALUES (?,?,?)",
                        (job_id, idx, '["user|pass"]'),
                    )
                claimed = set()
                for pool, expected in (("vvip", (count * 3 + 9) // 10),
                                       ("normal", count - (count * 3 + 9) // 10)):
                    pool_count = 0
                    while True:
                        row = master_server.select_fair_claim_candidate(
                            self.inner, time.time(), "test-satellite", pool
                        )
                        if row is None:
                            break
                        self.assertEqual(row[1], job_id)
                        self.assertNotIn(row[0], claimed)
                        claimed.add(row[0])
                        pool_count += 1
                        self.inner.exec("UPDATE chunks SET status='done' WHERE id=?", (row[0],))
                    self.assertEqual(pool_count, expected)
                self.assertEqual(len(claimed), count)
                self.inner.exec("UPDATE jobs SET status='done' WHERE id=?", (job_id,))

    def test_vvip_drains_normal_share_without_stealing_active_chunks(self) -> None:
        job_id = self.inner.exec(
            "INSERT INTO jobs (created_at,total,chunk_size,status,queue_type) VALUES (?,?,?,?,?)",
            (time.time(), 10, 1, "open", "vvip"),
        )
        chunk_ids = {}
        for idx in range(10):
            chunk_ids[idx] = self.inner.exec(
                "INSERT INTO chunks (job_id,idx,account) VALUES (?,?,?)",
                (job_id, idx, '["user|pass"]'),
            )
        normal_job = self.inner.exec(
            "INSERT INTO jobs (created_at,total,chunk_size,status,queue_type) VALUES (?,?,?,?,?)",
            (time.time(), 1, 1, "open", "normal"),
        )
        self.inner.exec(
            "INSERT INTO chunks (job_id,idx,account) VALUES (?,?,?)",
            (normal_job, 0, '["normal|pass"]'),
        )
        self.inner.exec(
            "UPDATE chunks SET status='claimed',satellite_id=?,lease_until=? WHERE id=?",
            ("normal-worker", time.time() + 600, chunk_ids[1]),
        )
        claimed = []
        for _ in range(9):
            status, payload = self.post("/api/vvip/claim", {"satellite_id": "vvip-worker"})
            self.assertEqual(status, 200)
            self.assertIsNotNone(payload["claim"])
            self.assertEqual(payload["claim"]["job_id"], job_id)
            claimed.append(payload["claim"]["chunk_id"])
        self.assertEqual(set(claimed[:3]), {chunk_ids[i] for i in (0, 3, 6)})
        self.assertEqual(set(claimed), set(chunk_ids.values()) - {chunk_ids[1]})
        self.assertEqual(len(set(claimed)), 9)
        _, payload = self.post("/api/vvip/claim", {"satellite_id": "vvip-worker"})
        self.assertIsNone(payload["claim"])
        self.assertEqual(
            self.inner.fetchone("SELECT satellite_id FROM chunks WHERE id=?", (chunk_ids[1],))[0],
            "normal-worker",
        )
        # A normal worker's expired lease becomes available to VVIP on the next poll.
        self.inner.exec("UPDATE chunks SET lease_until=? WHERE id=?", (time.time() - 1, chunk_ids[1]))
        _, payload = self.post("/api/vvip/claim", {"satellite_id": "vvip-worker"})
        self.assertEqual(payload["claim"]["chunk_id"], chunk_ids[1])

    def test_open_job_without_chunks_is_not_completed(self) -> None:
        orphan_id = self.inner.exec(
            "INSERT INTO jobs (created_at, total, chunk_size, status) VALUES (?,?,?,?)",
            (1, 1, 15, "open"),
        )
        handler = object.__new__(MasterHandler)
        handler.server = self.server
        handler._check_finish_all_jobs(2)
        self.assertEqual(
            self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (orphan_id,))[0],
            "open",
        )

    def test_regular_key_gets_global_overview_but_only_its_job_list(self) -> None:
        owner_a_job = self.inner.exec(
            "INSERT INTO jobs (created_at, total, chunk_size, status, owner_hash, owner_preview) VALUES (?,?,?,?,?,?)",
            (1, 100, 15, "open", "owner-a", "key-a"),
        )
        owner_b_job = self.inner.exec(
            "INSERT INTO jobs (created_at, total, chunk_size, status, owner_hash, owner_preview) VALUES (?,?,?,?,?,?)",
            (2, 250, 15, "done", "owner-b", "key-b"),
        )
        self.inner.exec(
            "INSERT INTO results (chunk_id, job_id, account, row_json, reported_at) VALUES (?,?,?,?,?)",
            (101, owner_a_job, "account-a", '{"status":"OK"}', 3),
        )
        self.inner.exec(
            "INSERT INTO results (chunk_id, job_id, account, row_json, reported_at) VALUES (?,?,?,?,?)",
            (202, owner_b_job, "account-b", '{"status":"FAIL"}', 3),
        )
        handler = object.__new__(MasterHandler)
        handler.server = self.server
        captured: dict = {}

        def capture_json(status: int, payload: dict) -> None:
            captured["status"] = status
            captured["payload"] = payload

        handler._json = capture_json
        handler._handle_jobs_list({"owner_hash": "owner-a", "is_admin": False})

        self.assertEqual(captured["status"], 200)
        self.assertEqual([job["owner_preview"] for job in captured["payload"]["jobs"]], ["key-a"])
        self.assertEqual(captured["payload"]["overview"], {
            "total_jobs": 2,
            "running_jobs": 1,
            "done_jobs": 1,
            "total_accounts": 350,
            "processed_accounts": 2,
            "pending_accounts": 99,
            "normal_running_jobs": 1,
            "vvip_running_jobs": 0,
        })

    def test_overview_splits_active_jobs_by_queue_across_owners(self) -> None:
        for queue in ("normal", "vvip"):
            for status in ("creating", "open", "stopping", "done"):
                self.inner.exec(
                    "INSERT INTO jobs (created_at,total,chunk_size,status,queue_type,owner_hash) "
                    "VALUES (?,?,?,?,?,?)",
                    (1, 1, 1, status, queue, "other-owner"),
                )
        handler = object.__new__(MasterHandler)
        handler.server = self.server
        captured = {}
        handler._json = lambda status, payload: captured.update(payload)
        handler._handle_jobs_list({"owner_hash": "owner-a", "is_admin": False})
        self.assertEqual(captured["jobs"], [])
        self.assertEqual(captured["overview"]["running_jobs"], 6)
        self.assertEqual(captured["overview"]["normal_running_jobs"], 3)
        self.assertEqual(captured["overview"]["vvip_running_jobs"], 3)

    def test_jobs_list_uses_constant_number_of_database_reads(self) -> None:
        for index in range(12):
            job_id = self.inner.exec(
                "INSERT INTO jobs (created_at, total, chunk_size, status, owner_hash, owner_preview) VALUES (?,?,?,?,?,?)",
                (index, 1, 15, "done", "owner-a", "key-a"),
            )
            self.inner.exec(
                "INSERT INTO results (chunk_id, job_id, account, row_json, reported_at) VALUES (?,?,?,?,?)",
                (1000 + index, job_id, f"account-{index}", '{"status":"OK"}', index),
            )

        counting_store = _CountingStore(self.inner)
        handler = object.__new__(MasterHandler)
        handler.server = type("Server", (), {"store": counting_store})()
        captured: dict = {}
        handler._json = lambda status, payload: captured.update(status=status, payload=payload)

        handler._handle_jobs_list({"owner_hash": "owner-a", "is_admin": False})

        self.assertEqual(captured["status"], 200)
        self.assertEqual(len(captured["payload"]["jobs"]), 12)
        self.assertTrue(all(job["processed"] == 1 for job in captured["payload"]["jobs"]))
        self.assertEqual(counting_store.read_calls, 2)

    def test_job_stats_follow_result_upserts_and_deletes(self) -> None:
        job_id = self.inner.exec(
            "INSERT INTO jobs (created_at, total, chunk_size, status) VALUES (?,?,?,?)",
            (1, 2, 15, "open"),
        )
        chunk_id = self.inner.exec(
            "INSERT INTO chunks (job_id, idx, account) VALUES (?,?,?)",
            (job_id, 0, "[]"),
        )
        insert_sql = (
            "INSERT INTO results (chunk_id, job_id, account, row_json, reported_at) "
            "VALUES (?,?,?,?,?) ON CONFLICT(chunk_id, account) DO UPDATE SET "
            "row_json=excluded.row_json, reported_at=excluded.reported_at"
        )
        self.inner.exec(insert_sql, (chunk_id, job_id, "a", '{"status":"OK"}', 2))
        self.assertEqual(
            self.inner.fetchone(
                "SELECT result_count, ok_count, fail_count, uncheckable_count FROM job_stats WHERE job_id=?",
                (job_id,),
            ),
            (1, 1, 0, 0),
        )
        self.assertEqual(
            self.inner.fetchone(
                "SELECT category, level, result_count FROM job_result_stats WHERE job_id=?",
                (job_id,),
            ),
            ("LEVEL", 0, 1),
        )

        self.inner.exec(insert_sql, (chunk_id, job_id, "a", '{"status":"FAIL"}', 3))
        self.assertEqual(
            self.inner.fetchone(
                "SELECT result_count, ok_count, fail_count, uncheckable_count FROM job_stats WHERE job_id=?",
                (job_id,),
            ),
            (1, 0, 1, 0),
        )
        self.assertEqual(
            self.inner.fetchone(
                "SELECT category, level, result_count FROM job_result_stats WHERE job_id=?",
                (job_id,),
            ),
            ("FAIL", 0, 1),
        )

        self.inner.exec("DELETE FROM results WHERE chunk_id=? AND account=?", (chunk_id, "a"))
        self.assertEqual(
            self.inner.fetchone(
                "SELECT result_count, ok_count, fail_count, uncheckable_count FROM job_stats WHERE job_id=?",
                (job_id,),
            ),
            (0, 0, 0, 0),
        )

    def test_all_correct_password_categories_count_as_ok_for_billing(self) -> None:
        job_id = self.inner.exec(
            "INSERT INTO jobs (created_at, total, chunk_size, status) VALUES (?,?,?,?)",
            (1, 4, 15, "open"),
        )
        chunk_id = self.inner.exec(
            "INSERT INTO chunks (job_id, idx, account) VALUES (?,?,?)",
            (job_id, 0, "[]"),
        )
        rows = (
            ("du-lv", {"status": "OK", "level": "50", "player_status": "Bình thường"}),
            ("chua-dat", {"status": "OK", "level": "5", "player_status": "Bình thường"}),
            ("bi-khoa", {"status": "OK", "level": "30", "player_status": "Bị khóa"}),
            ("ctnv", {"status": "OK", "level": "Ctnv", "player_status": "Chưa tạo nhân vật"}),
        )
        for account, row in rows:
            self.inner.exec(
                "INSERT INTO results (chunk_id, job_id, account, row_json, reported_at) VALUES (?,?,?,?,?)",
                (chunk_id, job_id, account, json.dumps(row, ensure_ascii=False), 2),
            )

        self.assertEqual(
            self.inner.fetchone(
                "SELECT result_count, ok_count, fail_count, uncheckable_count FROM job_stats WHERE job_id=?",
                (job_id,),
            ),
            (4, 4, 0, 0),
        )

    def test_job_stats_follow_chunk_and_job_status_updates(self) -> None:
        job_id = self.inner.exec(
            "INSERT INTO jobs (created_at, total, chunk_size, status) VALUES (?,?,?,?)",
            (1, 10, 15, "creating"),
        )
        chunk_id = self.inner.exec(
            "INSERT INTO chunks (job_id, idx, account) VALUES (?,?,?)",
            (job_id, 0, "[]"),
        )
        self.inner.exec("UPDATE jobs SET status='open' WHERE id=?", (job_id,))
        self.inner.exec("UPDATE chunks SET status='claimed' WHERE id=?", (chunk_id,))
        self.inner.exec("UPDATE chunks SET status='done' WHERE id=?", (chunk_id,))
        self.assertEqual(
            self.inner.fetchone(
                "SELECT total_accounts, job_status, pending_chunks, claimed_chunks, done_chunks "
                "FROM job_stats WHERE job_id=?",
                (job_id,),
            ),
            (10, "open", 0, 0, 1),
        )

    def test_existing_database_is_backfilled_once(self) -> None:
        legacy_path = Path(self.temp_dir.name) / "legacy.db"
        connection = sqlite3.connect(legacy_path)
        connection.executescript("""
            CREATE TABLE jobs (
                id INTEGER PRIMARY KEY AUTOINCREMENT, created_at REAL NOT NULL,
                total INTEGER NOT NULL, chunk_size INTEGER NOT NULL,
                status TEXT NOT NULL DEFAULT 'open', finished_at REAL,
                owner_hash TEXT DEFAULT '', owner_preview TEXT DEFAULT ''
            );
            CREATE TABLE chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER NOT NULL,
                idx INTEGER NOT NULL, account TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending', satellite_id TEXT DEFAULT '',
                claimed_at REAL, lease_until REAL, reported_at REAL,
                retry_round INTEGER NOT NULL DEFAULT 1, avoid_satellite_id TEXT DEFAULT '',
                UNIQUE(job_id, idx)
            );
            CREATE TABLE results (
                id INTEGER PRIMARY KEY AUTOINCREMENT, chunk_id INTEGER NOT NULL,
                job_id INTEGER NOT NULL, account TEXT NOT NULL, row_json TEXT NOT NULL,
                reported_at REAL NOT NULL, UNIQUE(chunk_id, account)
            );
            CREATE TABLE app_settings (setting_key TEXT PRIMARY KEY, setting_value TEXT NOT NULL);
            INSERT INTO jobs (id, created_at, total, chunk_size, status) VALUES (7, 1, 2, 15, 'open');
            INSERT INTO chunks (id, job_id, idx, account, status) VALUES (9, 7, 0, '[]', 'done');
            INSERT INTO results (chunk_id, job_id, account, row_json, reported_at)
            VALUES (9, 7, 'a', '{"status":"OK","level":"15"}', 2);
        """)
        connection.commit()
        connection.close()

        migrated = LocalStore(legacy_path)
        try:
            self.assertEqual(
                migrated.fetchone(
                    "SELECT total_accounts, done_chunks, result_count, ok_count FROM job_stats WHERE job_id=7"
                ),
                (2, 1, 1, 1),
            )
            self.assertEqual(
                migrated.fetchone(
                    "SELECT category, level, result_count FROM job_result_stats WHERE job_id=7"
                ),
                ("LEVEL", 15, 1),
            )
            self.assertEqual(
                migrated.fetchone(
                    "SELECT setting_value FROM app_settings WHERE setting_key='job_stats_version'"
                )[0],
                "2",
            )
        finally:
            migrated._conn.close()

    def test_job_summary_uses_one_database_read(self) -> None:
        job_id = self.inner.exec(
            "INSERT INTO jobs (created_at, total, chunk_size, status, owner_hash, owner_preview) VALUES (?,?,?,?,?,?)",
            (1, 3, 15, "open", "owner-a", "key-a"),
        )
        chunk_ids = [
            self.inner.exec(
                "INSERT INTO chunks (job_id, idx, account, status) VALUES (?,?,?,?)",
                (job_id, index, "[]", status),
            )
            for index, status in enumerate(("pending", "claimed", "done"))
        ]
        self.inner.exec(
            "INSERT INTO results (chunk_id, job_id, account, row_json, reported_at) VALUES (?,?,?,?,?)",
            (chunk_ids[-1], job_id, "account-ok", '{"status":"OK"}', 2),
        )

        counting_store = _CountingStore(self.inner)
        handler = object.__new__(MasterHandler)
        handler.server = type("Server", (), {"store": counting_store})()
        captured: dict = {}
        handler._json = lambda status, payload: captured.update(status=status, payload=payload)

        handler._handle_job_summary(job_id, {"owner_hash": "owner-a", "is_admin": False})

        self.assertEqual(captured["status"], 200)
        self.assertEqual(captured["payload"]["chunks"], {"pending": 1, "claimed": 1, "done": 1})
        self.assertEqual(captured["payload"]["results"]["count"], 1)
        self.assertEqual(counting_store.read_calls, 1)

    def test_notice_html_and_css_can_be_saved(self) -> None:
        handler = object.__new__(MasterHandler)
        handler.server = self.server
        captured: list[tuple[int, dict]] = []
        notice_html = '<div class="notice-title">Bảo trì</div><div>Thông báo mới</div>'
        notice_css = '#noticeBox { background: #fff; }'
        handler._read_json = lambda: {
            "notice": {"enabled": True, "html": notice_html, "css": notice_css},
        }
        handler._json = lambda status, payload: captured.append((status, payload))

        handler._handle_admin_settings_save()
        self.assertEqual(captured[-1][0], 200)
        handler._handle_admin_settings_get()

        returned = captured[-1][1]["notice"]
        self.assertEqual(returned["html"], notice_html)
        self.assertEqual(returned["css"], notice_css)
        self.assertTrue(returned["enabled"])

    def test_admin_can_save_and_read_running_job_limit(self) -> None:
        handler = object.__new__(MasterHandler)
        handler.server = self.server
        captured: list[tuple[int, dict]] = []
        handler._read_json = lambda: {"max_running_jobs": 3}
        handler._json = lambda status, payload: captured.append((status, payload))

        handler._handle_admin_settings_save()
        self.assertEqual(captured[-1][0], 200)
        handler._handle_admin_settings_get()

        settings = captured[-1][1]
        self.assertEqual(settings["max_running_jobs"], 3)
        self.assertEqual(settings["running_jobs"], 0)
        self.assertEqual(settings["available_job_slots"], 3)

    def test_admin_can_save_and_read_account_limit_per_job(self) -> None:
        handler = object.__new__(MasterHandler)
        handler.server = self.server
        captured: list[tuple[int, dict]] = []
        handler._read_json = lambda: {"max_accounts_per_job": 2500}
        handler._json = lambda status, payload: captured.append((status, payload))

        handler._handle_admin_settings_save()
        self.assertEqual(captured[-1][0], 200)
        handler._handle_admin_settings_get()

        self.assertEqual(captured[-1][1]["max_accounts_per_job"], 2500)

    def test_account_limit_rejects_regular_key_but_not_admin_key(self) -> None:
        self.store.block_once = False
        self.inner.exec(
            "INSERT INTO app_settings (setting_key, setting_value) VALUES (?,?)",
            ("max_accounts_per_job", "1"),
        )
        text = "user01|password1!\nuser02|password2!"

        status, rejected = self.post_error("/api/jobs", {"text": text}, token="regular-key")
        self.assertEqual(status, 413)
        self.assertEqual(rejected["code"], "ACCOUNT_LIMIT_REACHED")
        self.assertEqual(rejected["submitted_accounts"], 2)
        self.assertEqual(rejected["max_accounts_per_job"], 1)
        self.assertEqual(self.inner.fetchone("SELECT COUNT(*) FROM jobs")[0], 0)

        status, created = self.post("/api/jobs", {"text": text}, token="secret")
        self.assertEqual(status, 200)
        self.assertEqual(created["total"], 2)

    def test_each_regular_key_can_only_have_one_running_job(self) -> None:
        self.store.block_once = False

        status, first = self.post("/api/jobs", {"text": "user01|password1!"}, token="key-a")
        self.assertEqual(status, 200)

        status, rejected = self.post_error("/api/jobs", {"text": "user02|password2!"}, token="key-a")
        self.assertEqual(status, 409)
        self.assertEqual(rejected["code"], "KEY_RUNNING_JOB_LIMIT_REACHED")
        self.assertEqual(rejected["active_job_id"], first["job_id"])

        status, other_key_job = self.post("/api/jobs", {"text": "user03|password3!"}, token="key-b")
        self.assertEqual(status, 200)
        self.assertNotEqual(other_key_job["job_id"], first["job_id"])

    def test_expired_key_can_read_history_but_cannot_create_job(self) -> None:
        self.store.block_once = False
        _, created = self.post("/api/jobs", {"text": "history-user|password!"}, token="expired-key")
        job_id = created["job_id"]

        with mock.patch.object(
            master_server,
            "_verify_license_key",
            return_value=(False, {"error": "key expired"}),
        ):
            status, verification = self.get("/api/verify", token="expired-key")
            self.assertEqual(status, 200)
            self.assertFalse(verification["valid"])
            self.assertTrue(verification["history_access"])
            self.assertFalse(verification["can_create_job"])

            status, history = self.get("/api/jobs_list", token="expired-key")
            self.assertEqual(status, 200)
            self.assertEqual([job["id"] for job in history["jobs"]], [job_id])

            status, summary = self.get(f"/api/jobs/{job_id}", token="expired-key")
            self.assertEqual(status, 200)
            self.assertEqual(summary["job_id"], job_id)

            status, stopped = self.post(f"/api/jobs/{job_id}/stop", {}, token="expired-key")
            self.assertEqual(status, 202)
            self.assertEqual(stopped["status"], "stopping")
            self.wait_for_job_status(job_id)

            status, rows = self.get(f"/api/jobs/{job_id}/rows", token="expired-key")
            self.assertEqual(status, 200)
            self.assertEqual(rows["total"], 1)
            self.assertEqual(len(rows["rows"]), 1)

            status, forbidden = self.get_error(f"/api/jobs/{job_id}", token="different-key")
            self.assertEqual(status, 403)
            self.assertFalse(forbidden["ok"])

            status, rejected = self.post_error(
                "/api/jobs", {"text": "new-user|password!"}, token="expired-key"
            )
            self.assertEqual(status, 401)
            self.assertFalse(rejected["ok"])

    def test_admin_key_is_exempt_from_running_job_limit_per_key(self) -> None:
        self.store.block_once = False

        status, first = self.post("/api/jobs", {"text": "admin1|password1!"})
        self.assertEqual(status, 200)

        status, second = self.post("/api/jobs", {"text": "admin2|password2!"})
        self.assertEqual(status, 200)
        self.assertNotEqual(second["job_id"], first["job_id"])

    def test_retention_keeps_old_running_job_until_it_finishes(self) -> None:
        cutoff = _today_start_timestamp()

        def add_job(created_at: float, status: str, account: str, chunk_status: str) -> tuple[int, int]:
            job_id = self.inner.exec(
                "INSERT INTO jobs (created_at, total, chunk_size, status, finished_at) VALUES (?,?,?,?,?)",
                (created_at, 1, 15, status, created_at if status == "done" else None),
            )
            chunk_id = self.inner.exec(
                "INSERT INTO chunks (job_id, idx, account, status) VALUES (?,?,?,?)",
                (job_id, 0, json.dumps([f"{account}|pass"]), chunk_status),
            )
            self.inner.exec(
                "INSERT INTO results (chunk_id, job_id, account, row_json, reported_at) VALUES (?,?,?,?,?)",
                (chunk_id, job_id, account, '{}', created_at),
            )
            return job_id, chunk_id

        old_done, _ = add_job(cutoff - 7200, "done", "old-done", "done")
        old_running, old_running_chunk = add_job(cutoff - 3600, "open", "old-running", "pending")
        today_done, _ = add_job(cutoff + 60, "done", "today-done", "done")

        deleted = _prune_completed_jobs_before_today(self.inner, cutoff)
        self.assertEqual(deleted, {"jobs": 1, "chunks": 1, "results": 1})
        self.assertIsNone(self.inner.fetchone("SELECT id FROM jobs WHERE id=?", (old_done,)))
        self.assertEqual(self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (old_running,))[0], "open")
        self.assertIsNotNone(self.inner.fetchone("SELECT id FROM jobs WHERE id=?", (today_done,)))

        self.inner.exec("UPDATE chunks SET status='done' WHERE id=?", (old_running_chunk,))
        handler = object.__new__(MasterHandler)
        handler.server = self.server
        handler._check_finish_all_jobs(cutoff + 120)

        self.assertIsNone(self.inner.fetchone("SELECT id FROM jobs WHERE id=?", (old_running,)))
        self.assertIsNotNone(self.inner.fetchone("SELECT id FROM jobs WHERE id=?", (today_done,)))

    def test_rejects_new_job_at_limit_and_accepts_after_a_job_stops(self) -> None:
        # This test exercises normal creation; the blocking wrapper is only needed
        # by the separate race test above.
        self.store.block_once = False
        self.inner.exec(
            "INSERT INTO app_settings (setting_key, setting_value) VALUES (?,?)",
            ("max_running_jobs", "1"),
        )

        status, first = self.post("/api/jobs", {"text": "user01|password1!"}, token="key-a")
        self.assertEqual(status, 200)

        status, rejected = self.post_error("/api/jobs", {"text": "user02|password2!"}, token="key-b")
        self.assertEqual(status, 429)
        self.assertEqual(rejected["code"], "JOB_LIMIT_REACHED")
        self.assertEqual(rejected["running_jobs"], 1)
        self.assertEqual(rejected["max_running_jobs"], 1)
        self.assertEqual(self.inner.fetchone("SELECT COUNT(*) FROM jobs")[0], 1)

        status, stopped = self.post(f"/api/jobs/{first['job_id']}/stop", {})
        self.assertEqual(status, 202)
        self.assertEqual(stopped["status"], "stopping")
        self.wait_for_job_status(first["job_id"])

        status, second = self.post("/api/jobs", {"text": "user02|password2!"}, token="key-b")
        self.assertEqual(status, 200)
        self.assertNotEqual(second["job_id"], first["job_id"])

    def test_two_jobs_can_be_stopped_at_the_same_time(self) -> None:
        self.store.block_once = False
        _, first = self.post("/api/jobs", {"text": "user01|password1!\nuser02|password2!"}, token="key-a")
        _, second = self.post("/api/jobs", {"text": "user03|password3!\nuser04|password4!"}, token="key-b")

        original = MasterHandler._finalize_unresolved_accounts
        counter_lock = threading.Lock()
        active = 0
        max_active = 0

        def measured_finalize(handler, job_id: int, now: float) -> int:
            nonlocal active, max_active
            with counter_lock:
                active += 1
                max_active = max(max_active, active)
            try:
                # Keep the first finalization open long enough for the other
                # request thread to reach the server.
                time.sleep(0.1)
                return original(handler, job_id, now)
            finally:
                with counter_lock:
                    active -= 1

        responses: list[tuple[int, dict]] = []
        errors: list[Exception] = []

        def stop(job_id: int, token: str) -> None:
            try:
                responses.append(self.post(f"/api/jobs/{job_id}/stop", {}, token=token))
            except Exception as exc:  # pragma: no cover - surfaced below
                errors.append(exc)

        with mock.patch.object(MasterHandler, "_finalize_unresolved_accounts", measured_finalize):
            threads = [
                threading.Thread(target=stop, args=(first["job_id"], "key-a")),
                threading.Thread(target=stop, args=(second["job_id"], "key-b")),
            ]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join(10)
            for job_id in (first["job_id"], second["job_id"]):
                self.wait_for_job_status(job_id)

        self.assertFalse(errors)
        self.assertTrue(all(not thread.is_alive() for thread in threads))
        self.assertEqual(sorted(status for status, _ in responses), [202, 202])
        self.assertEqual(max_active, 1)
        for job_id in (first["job_id"], second["job_id"]):
            self.assertEqual(self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (job_id,))[0], "done")
            self.assertEqual(
                self.inner.fetchone("SELECT COUNT(*) FROM chunks WHERE job_id=? AND status!='done'", (job_id,))[0],
                0,
            )
            self.assertEqual(
                self.inner.fetchone("SELECT COUNT(*) FROM results WHERE job_id=?", (job_id,))[0],
                2,
            )

    def test_stop_request_returns_before_slow_finalization_finishes(self) -> None:
        self.store.block_once = False
        _, created = self.post(
            "/api/jobs", {"text": "user01|password1!\nuser02|password2!"}, token="key-a"
        )
        job_id = created["job_id"]
        entered = threading.Event()
        release = threading.Event()
        original = MasterHandler._finalize_unresolved_accounts

        def blocking_finalize(handler, current_job_id: int, now: float) -> int:
            entered.set()
            if not release.wait(5):
                raise TimeoutError("test did not release stop finalization")
            return original(handler, current_job_id, now)

        with mock.patch.object(
            MasterHandler, "_finalize_unresolved_accounts", blocking_finalize
        ):
            started = time.monotonic()
            status, stopped = self.post(f"/api/jobs/{job_id}/stop", {}, token="key-a")
            elapsed = time.monotonic() - started
            self.assertEqual(status, 202)
            self.assertEqual(stopped["status"], "stopping")
            self.assertLess(elapsed, 1.0)
            self.assertTrue(entered.wait(2))
            self.assertEqual(
                self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (job_id,))[0],
                "stopping",
            )
            release.set()
            self.wait_for_job_status(job_id)
            self.assertEqual(self.inner.fetchone("SELECT COUNT(*) FROM results WHERE job_id=?", (job_id,))[0], 2)

    def test_quantity_settlement_retries_and_is_retained_until_confirmed(self) -> None:
        cutoff = _today_start_timestamp()
        job_id = self.inner.exec(
            "INSERT INTO jobs (created_at,total,chunk_size,status,finished_at,owner_user_id,"
            "billing_mode,billing_state,external_job_reference,unit_price_tenths,estimated_amount_tenths) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (cutoff - 3600, 2, 15, "done", cutoff - 3500, 77, "quantity", "reserved", "cp-test-settle", 3, 6),
        )
        chunk_id = self.inner.exec(
            "INSERT INTO chunks (job_id,idx,account,status) VALUES (?,?,?,?)",
            (job_id, 0, json.dumps(["ok|pass", "bad|pass"]), "done"),
        )
        for account, status in (("ok", "OK"), ("bad", "FAIL")):
            self.inner.exec(
                "INSERT INTO results (chunk_id,job_id,account,row_json,reported_at) VALUES (?,?,?,?,?)",
                (chunk_id, job_id, account, json.dumps({"status": status}), cutoff - 3500),
            )

        with mock.patch.object(master_server, "_aovshop_configured", return_value=True), mock.patch.object(
            master_server, "_aovshop_request", side_effect=RuntimeError("temporary outage")
        ):
            self.server._settle_billing_job(job_id)

        state = self.inner.fetchone("SELECT billing_state,billing_error FROM jobs WHERE id=?", (job_id,))
        self.assertEqual(state[0], "settlement_pending")
        self.assertIn("temporary outage", state[1])
        self.assertIsNotNone(self.inner.fetchone("SELECT job_id FROM billing_outbox WHERE job_id=?", (job_id,)))
        self.assertEqual(_prune_completed_jobs_before_today(self.inner, cutoff)["jobs"], 0)
        # Keep the row visible after successful settlement so its persisted
        # fields can be asserted; pruning settled old rows is covered below.
        self.inner.exec("UPDATE jobs SET created_at=? WHERE id=?", (cutoff + 1, job_id))

        captured: dict = {}

        def settle(path: str, payload: dict, method: str = "POST") -> dict:
            captured.update(payload)
            return {"ok": True, "status": "settled", "order_id": 501}

        with mock.patch.object(master_server, "_aovshop_configured", return_value=True), mock.patch.object(
            master_server, "_aovshop_request", side_effect=settle
        ):
            self.server._settle_billing_job(job_id)

        self.assertEqual(captured["ok_count"], 1)
        self.assertEqual(captured["fail_count"], 1)
        self.assertEqual(captured["uncheckable_count"], 0)
        self.assertEqual(captured["idempotency_key"], "settle:cp-test-settle")
        settled = self.inner.fetchone(
            "SELECT billing_state,final_amount_tenths,billing_order_id FROM jobs WHERE id=?", (job_id,)
        )
        self.assertEqual(settled, ("settled", 4, 501))
        self.assertIsNone(self.inner.fetchone("SELECT job_id FROM billing_outbox WHERE job_id=?", (job_id,)))
        self.inner.exec("UPDATE jobs SET created_at=? WHERE id=?", (cutoff - 1, job_id))
        self.assertEqual(_prune_completed_jobs_before_today(self.inner, cutoff)["jobs"], 1)

    def test_new_job_resets_satellite_idle_restart_timer(self) -> None:
        self.store.block_once = False
        self.server._satellite_idle_since = 0
        self.server._satellite_idle_restarted = True
        status, created = self.post("/api/jobs", {"text": "idle-reset|password!"})
        self.assertEqual(status, 200)
        self.assertTrue(created["job_id"])
        self.assertIsNone(self.server._satellite_idle_since)
        self.assertFalse(self.server._satellite_idle_restarted)

    def test_expired_time_job_remains_available_for_new_claims(self) -> None:
        now = time.time()
        expired_id = self.inner.exec(
            "INSERT INTO jobs (created_at,total,chunk_size,status,billing_mode,billing_state,access_until) "
            "VALUES (?,?,?,?,?,?,?)",
            (now - 100, 1, 15, "open", "time", "settled", now - 1),
        )
        self.inner.exec(
            "INSERT INTO chunks (job_id,idx,account,status) VALUES (?,?,?,?)",
            (expired_id, 0, json.dumps(["expired|pass"]), "pending"),
        )
        active_id = self.inner.exec(
            "INSERT INTO jobs (created_at,total,chunk_size,status,billing_mode,billing_state,access_until) "
            "VALUES (?,?,?,?,?,?,?)",
            (now, 1, 15, "open", "time", "settled", now + 60),
        )
        self.inner.exec(
            "INSERT INTO chunks (job_id,idx,account,status) VALUES (?,?,?,?)",
            (active_id, 0, json.dumps(["active|pass"]), "pending"),
        )

        claim = master_server.select_fair_claim_candidate(self.inner, now, "satellite-a")
        self.assertIsNotNone(claim)
        self.assertEqual(claim[1], expired_id)

    def test_billing_maintenance_keeps_expired_rental_jobs_open(self) -> None:
        now = time.time()
        job_ids = []
        for mode in ("time", "vvip"):
            job_ids.append(self.inner.exec(
                "INSERT INTO jobs (created_at,total,chunk_size,status,billing_mode,billing_state,access_until) "
                "VALUES (?,?,?,?,?,?,?)",
                (now - 100, 1, 15, "open", mode, "settled", now - 1),
            ))

        stop_event = threading.Event()
        with mock.patch.object(stop_event, "wait", side_effect=lambda _: stop_event.set()), mock.patch.object(
            self.server, "schedule_stop_finalization"
        ) as finalize:
            self.server.billing_maintenance_loop(stop_event)
        finalize.assert_not_called()
        for job_id in job_ids:
            self.assertEqual(
                self.inner.fetchone("SELECT status FROM jobs WHERE id=?", (job_id,))[0], "open"
            )
            self.inner.exec(
                "INSERT INTO chunks (job_id,idx,account,status) VALUES (?,?,?,?)",
                (job_id, 0, json.dumps(["expired|pass"]), "pending"),
            )
            self.inner.exec("UPDATE jobs SET queue_type=billing_mode WHERE id=? AND billing_mode='vvip'", (job_id,))
            queue = "vvip" if job_id == job_ids[1] else "normal"
            claim = master_server.select_fair_claim_candidate(self.inner, now, "satellite-a", queue)
            self.assertIsNotNone(claim)
            self.assertEqual(claim[1], job_id)

    def test_sp1s_mode_never_falls_back_to_open_access_without_master_token(self) -> None:
        original_token = self.server.master_token
        self.server.master_token = ""
        try:
            with mock.patch.object(master_server, "_aovshop_configured", return_value=True):
                status, user_error = self.get_error("/api/jobs_list", token="")
                self.assertEqual(status, 401)
                self.assertIn("SP1S", user_error["error"])

                status, satellite_error = self.post_error(
                    "/api/claim", {"satellite_id": "unauthenticated"}, token=""
                )
                self.assertEqual(status, 503)
                self.assertIn("MASTER_TOKEN", satellite_error["error"])
        finally:
            self.server.master_token = original_token

    def test_sp1s_login_uses_state_and_creates_http_only_session(self) -> None:
        class NoRedirect(urllib.request.HTTPRedirectHandler):
            def redirect_request(self, req, fp, code, msg, headers, newurl):
                return None

        opener = urllib.request.build_opener(NoRedirect())

        def redirect(path: str, cookie: str = ""):
            headers = {"Cookie": cookie} if cookie else {}
            try:
                opener.open(urllib.request.Request(self.base_url + path, headers=headers), timeout=5)
            except urllib.error.HTTPError as exc:
                self.assertEqual(exc.code, 303)
                return exc.headers
            self.fail("expected redirect")

        def aov_response(path: str, payload=None, method: str = "POST") -> dict:
            if path.endswith("/sso/exchange"):
                return {"ok": True, "user": {"id": 91, "name": "SP1S User", "email": "user@example.test"}}
            if "/account/91" in path:
                return {
                    "ok": True,
                    "user": {"id": 91, "name": "SP1S User", "email": "user@example.test"},
                    "balance": 9999.4,
                    "available_balance": 9999.4,
                    "entitlement": None,
                }
            raise AssertionError(path)

        with mock.patch.object(master_server, "_aovshop_configured", return_value=True), mock.patch.object(
            master_server, "_aovshop_request", side_effect=aov_response
        ):
            register_headers = redirect("/auth/register")
            register_location = urllib.parse.urlparse(register_headers["Location"])
            self.assertEqual(register_location.path, "/register")
            register_redirect = urllib.parse.parse_qs(register_location.query)["redirect"][0]
            register_connect = urllib.parse.urlparse(register_redirect)
            self.assertEqual(register_connect.path, "/checkpass/connect")
            register_return_url = urllib.parse.parse_qs(register_connect.query)["return_url"][0]
            self.assertEqual(urllib.parse.urlparse(register_return_url).path, "/auth/callback")

            login_headers = redirect("/auth/login")
            location = login_headers["Location"]
            login_location = urllib.parse.urlparse(location)
            self.assertEqual(login_location.path, "/login")
            login_redirect = urllib.parse.parse_qs(login_location.query)["redirect"][0]
            login_connect = urllib.parse.urlparse(login_redirect)
            self.assertEqual(login_connect.path, "/checkpass/connect")
            return_url = urllib.parse.parse_qs(login_connect.query)["return_url"][0]
            state = urllib.parse.parse_qs(urllib.parse.urlparse(return_url).query)["state"][0]
            state_cookie = login_headers.get_all("Set-Cookie")[0].split(";", 1)[0]
            self.assertIn("HttpOnly", login_headers.get_all("Set-Cookie")[0])

            callback_headers = redirect(
                "/auth/callback?state=" + urllib.parse.quote(state) + "&code=" + "a" * 32,
                state_cookie,
            )
            session_header = next(
                value for value in callback_headers.get_all("Set-Cookie")
                if value.startswith("checkpass_session=") and "Max-Age=0" not in value
            )
            self.assertIn("HttpOnly", session_header)
            session_cookie = session_header.split(";", 1)[0]
            request = urllib.request.Request(self.base_url + "/api/session", headers={"Cookie": session_cookie})
            with urllib.request.urlopen(request, timeout=5) as response:
                payload = json.loads(response.read().decode("utf-8"))
            self.assertEqual(payload["user"]["id"], 91)
            self.assertEqual(payload["balance"], 9999.4)


if __name__ == "__main__":
    unittest.main()
